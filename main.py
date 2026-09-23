import vk_api
from vk_api.utils import get_random_id
import re
import os
import json
import time
import tempfile
import threading
import asyncio
import logging
import psycopg2
from psycopg2.extras import RealDictCursor
from psycopg2 import pool as psycopg2_pool
from datetime import datetime, date
from fastapi import FastAPI, Request, HTTPException
from fastapi.responses import PlainTextResponse
import requests

# --- ЛОГИРОВАНИЕ ---
logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("vk_bot")

# --- РЕГЕКСЫ ---
PHONE_PATTERN = re.compile(r'^(\+7|7|8)?[\s\-]?\(?\d{3}\)?[\s\-]?\d{3}[\s\-]?\d{2}[\s\-]?\d{2}$')
EMAIL_PATTERN = re.compile(r'^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$')
NAME_PATTERN = re.compile(r"^[а-яёА-ЯЁa-zA-Z]+(?:['\-][а-яёА-ЯЁa-zA-Z]+)*$")

# --- ВСПОМОГАТЕЛЬНЫЕ ---
def format_numbered_list(items, start_from=1, truncate=True):
    lines = []
    for i, item in enumerate(items, start=start_from):
        if truncate and len(item) > 80:
            item = item[:77] + "…"
        lines.append(f"{i} — {item}")
    return "\n".join(lines)


def validate_fio(text):
    text = text.strip()
    lower = text.lower()
    if lower in ("начать анкету", "пройти заново", "🔄 пройти заново", "/restart"):
        return False, (
            "Кажется, вы нажали кнопку вместо ответа 🙂\n"
            "Пожалуйста, укажите Фамилию Имя и Отчество через пробел.\n"
            "Например: Иванов Иван Иванович\n"
            "Если отчества нет — поставьте «-» (например: Иванов Иван -)."
        )
    parts = text.split()
    if len(parts) != 3:
        return False, (
            "Нужно указать ровно три слова: Фамилия, Имя, Отчество — через пробел.\n"
            "Например: Иванов Иван Иванович\n"
            "Если отчества нет — поставьте «-» (например: Иванов Иван -)."
        )
    surname, first_name, patronymic = parts
    errors = []
    if not NAME_PATTERN.match(surname) or len(surname) < 2:
        errors.append("фамилия")
    if not NAME_PATTERN.match(first_name) or len(first_name) < 2:
        errors.append("имя")
    if patronymic != "-" and (not NAME_PATTERN.match(patronymic) or len(patronymic) < 2):
        errors.append("отчество")
    if errors:
        hint = (
            "Проверьте следующие поля: " + ", ".join(errors) + ".\n"
            "Используйте только буквы (допускается дефис в двойных фамилиях).\n"
            "Если отчества нет — поставьте «-» (например: Иванов Иван -)."
        )
        return False, hint

    def cap(s):
        if s == "-":
            return s
        return "-".join(p.capitalize() for p in s.split("-"))

    normalized = f"{cap(surname)} {cap(first_name)} {cap(patronymic)}"
    return True, normalized


def validate_contacts(text):
    text = text.strip()
    parts = re.split(r'[\s,;]+', text)
    phone = None
    email = None
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if phone is None and PHONE_PATTERN.match(part):
            digits = re.sub(r'\D', '', part)
            if len(digits) == 11 and digits.startswith('8'):
                digits = '7' + digits[1:]
            elif len(digits) == 10:
                digits = '7' + digits
            if len(digits) == 11:
                phone = '+7' + digits[-10:]
        elif email is None and EMAIL_PATTERN.match(part):
            email = part.lower()
    if phone or email:
        contacts = []
        if phone:
            contacts.append(phone)
        if email:
            contacts.append(email)
        return True, "; ".join(contacts)
    return False, None


# ================= НАСТРОЙКИ ИЗ ENV =================
VK_TOKEN = os.getenv("VK_TOKEN")
ADMIN_IDS = set(map(int, os.getenv("ADMIN_IDS", "").split(","))) if os.getenv("ADMIN_IDS") else set()
GROUP_ID = int(os.getenv("GROUP_ID"))
DATABASE_URL = os.getenv("DATABASE_URL")
CALLBACK_CONFIRM_TOKEN = os.getenv("CALLBACK_CONFIRM_TOKEN", "")

if not VK_TOKEN or not DATABASE_URL:
    raise ValueError("Не заданы переменные окружения VK_TOKEN или DATABASE_URL")
# =====================================================
vk_session = vk_api.VkApi(token=VK_TOKEN)
vk = vk_session.get_api()

# ----------------- ПУЛ СОЕДИНЕНИЙ БД -----------------
db_pool = None

def init_db_pool():
    global db_pool
    db_pool = psycopg2_pool.ThreadedConnectionPool(
        minconn=2,
        maxconn=30,
        dsn=DATABASE_URL
    )

def get_db():
    return db_pool.getconn()

def release_db(conn):
    db_pool.putconn(conn)

def init_db():
    conn = get_db()
    try:
        c = conn.cursor()
        c.execute("""CREATE TABLE IF NOT EXISTS answers (
                user_id BIGINT PRIMARY KEY,
                fio TEXT, institution TEXT, specialty TEXT, study_group TEXT,
                course TEXT, form_of_study TEXT, contacts TEXT,
                employment_status TEXT, target_contract TEXT, experience TEXT,
                practice_eval TEXT, events TEXT, resume_status TEXT,
                interview_training TEXT, special_status TEXT, military TEXT,
                maternity TEXT, graduate TEXT, post_plans TEXT, help_needed TEXT)""")
        c.execute("ALTER TABLE answers ADD COLUMN IF NOT EXISTS consent_status BOOLEAN")
        c.execute("ALTER TABLE answers ADD COLUMN IF NOT EXISTS created_at TIMESTAMP")
        c.execute("""CREATE TABLE IF NOT EXISTS progress (
                user_id BIGINT PRIMARY KEY,
                step_index INTEGER DEFAULT 0,
                uni_page INTEGER DEFAULT 0,
                started INTEGER DEFAULT 0)""")
        c.execute("ALTER TABLE progress ADD COLUMN IF NOT EXISTS started_at TIMESTAMP")
        c.execute("CREATE INDEX IF NOT EXISTS idx_progress_user_id ON progress(user_id)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_answers_user_id ON answers(user_id)")
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        release_db(conn)


def get_progress(user_id):
    conn = get_db()
    try:
        c = conn.cursor()
        c.execute("SELECT step_index, uni_page, started FROM progress WHERE user_id=%s", (user_id,))
        row = c.fetchone()
        if row:
            step_index = row[0] if row[0] is not None else 0
            uni_page = row[1] if row[1] is not None else 0
            started = row[2] if row[2] is not None else 0
            return step_index, uni_page, started
        return 0, 0, 0
    finally:
        release_db(conn)

def set_progress(user_id, step_index, uni_page=0, started=1):
    conn = get_db()
    try:
        c = conn.cursor()
        if started == 1:
            c.execute(
                "INSERT INTO progress (user_id, step_index, uni_page, started, started_at) "
                "VALUES (%s,%s,%s,%s,%s) "
                "ON CONFLICT (user_id) DO UPDATE SET step_index=%s, uni_page=%s, started=%s, "
                "started_at=COALESCE(progress.started_at, EXCLUDED.started_at)",
                (user_id, step_index, uni_page, started, datetime.now(),
                 step_index, uni_page, started)
            )
        else:
            c.execute(
                "INSERT INTO progress (user_id, step_index, uni_page, started) "
                "VALUES (%s,%s,%s,%s) "
                "ON CONFLICT (user_id) DO UPDATE SET step_index=%s, uni_page=%s, started=%s",
                (user_id, step_index, uni_page, started, step_index, uni_page, started)
            )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        release_db(conn)

def save_answer(user_id, field, value):
    cols = ["fio","institution","specialty","study_group","course","form_of_study",
            "contacts","employment_status","target_contract","experience",
            "practice_eval","events","resume_status","interview_training",
            "special_status","military","maternity","graduate","post_plans","help_needed"]
    if field not in cols:
        return
    conn = get_db()
    try:
        c = conn.cursor()
        c.execute(
            f"INSERT INTO answers (user_id, {field}) VALUES (%s, %s) "
            f"ON CONFLICT (user_id) DO UPDATE SET {field}=EXCLUDED.{field}",
            (user_id, value)
        )
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        release_db(conn)

def check_answered(user_id, step_key):
    conn = get_db()
    try:
        c = conn.cursor()
        if step_key == "consent":
            c.execute("SELECT consent_status FROM answers WHERE user_id=%s", (user_id,))
            row = c.fetchone()
            return row is not None and row[0] is True
        elif step_key == "fio":
            c.execute("SELECT fio FROM answers WHERE user_id=%s", (user_id,))
            row = c.fetchone()
            return row is not None and row[0]
        else:
            field = STEP_TO_DB.get(step_key, step_key)
            c.execute(f"SELECT {field} FROM answers WHERE user_id=%s", (user_id,))
            row = c.fetchone()
            return row is not None and row[0]
    finally:
        release_db(conn)

# ----------------- ВУЗы -----------------
UNIVERSITIES = [
    "БПОУ УР «Воткинский промышленный техникум»",
    "БПОУ УР «Воткинский музыкально-педагогический колледж имени П.И. Чайковского»",
    "БПОУ УР «Воткинский машиностроительный техникум имени В.Г.Садовникова»",
    "АПОУ УР «Глазовский аграрно-промышленный техникум»",
    "БПОУ УР «Глазовский технический колледж»",
    "БПОУ «Дебесский политехникум»",
    "БПОУР «Игринский политехнический техникум»",
    "БПОУ УР «Ижевский торгово-экономический техникум»",
    "БПОУ УР «Ижевский монтажный техникум»",
    "БПОУ «Ижевский агростроительный техникум»",
    "ЧПОО «Нефтяной техникум»",
    "БПОУ УР «Ижевский политехнический колледж»",
    "БПОУ УР «Ижевский промышленно-экономический колледж»",
    "БПОУ УР «Ижевский машиностроительный техникум им. С.Н. Борина»",
    "БПОУ УР «Радиомеханический техникум имени В.А. Шутова»",
    "АПОУ УР «Экономико-технологический колледж»",
    "БПОУ УР «Асановский аграрно-технический техникум»",
    "АПОУ УР «Топливно-энергетический колледж»",
    "БПОУ УР «Ижевский техникум индустрии питания»",
    "КПОУ УР «Удмуртский республиканский колледж культуры»",
    "АНПОО «Международный Восточно-Европейский колледж»",
    "АПОУ УР «Техникум радиоэлектроники и информационных технологий им. А.В. Воскресенского»",
    "АПОУ УР «Республиканский медицинский колледж имени Героя Советского Союза Ф.А. Пушиной Министерства здравоохранения Удмуртской Республики»",
    "Республиканский медицинский колледж имени Героя Советского Союза Ф.А. Пушиной, Глазовский филиал",
    "Воткинский филиал АПОУ УР «Республиканский медицинский колледж имени Героя Советского Союза Ф.А. Пушиной Министерства здравоохранения Удмуртской Республики»",
    "ПОЧУ «Ижевский техникум экономики, управления и права Удмуртпотребсоюза»",
    "АНПОО СПО «Ижевский финансово-юридический колледж»",
    "БПОУ УР «Удмуртский республиканский социально-педагогический колледж»",
    "АПОУ УР «Строительный техникум»",
    "ФГБОУ ВО «Ижевская государственная медицинская академия»",
    "КПОУ УР «Республиканский музыкальный колледж»",
    "БПОУ УР «Ижевский промышленно-экономический колледж» в г. Можга",
    "БПОУ УР «Можгинский педагогический колледж имени Т.К. Борисова»",
    "БПОУ УР «Можгинский агропромышленный колледж»",
    "БПОУ УР «Сарапульский политехнический техникум»",
    "БПОУ УР «Сарапульский многопрофильный колледж»",
    "БПОУ УР «Сарапульский колледж социально-педагогических технологий и сервиса»",
    "БПОУ УР «Сюмсинский техникум лесного и сельского хозяйства»",
    "БПОУ УР «Увинский профессиональный колледж»",
    "БПОУ УР «Ярский политехникум»",
    "ФГБОУ ВО «Приволжский государственный университет путей сообщения»",
    "БПОУ УР «Ижевский индустриальный техникум имени Евгения Фёдоровича Драгунова»",
    "БПОУ УР «Глазовский политехнический колледж»",
    "Ижевский институт (филиал) ВГУЮ (РПА Минюста России)",
    "ФГБОУ ВО «Удмуртский государственный университет»",
    "ФГБОУ ВО «Удмуртский государственный аграрный университет»",
    "ФГБОУ ВО «Ижевский государственный технический университет имени М.Т. Калашникова»",
    "Глазовский инженерно-экономический институт(филиал) ФГБОУ ВО «ИжГТУ имени М.Т. Калашникова»",
    "БПОУ УР «Ижевский автотранспортный техникум»",
    "ФГБОУ ВО «Глазовский государственный инженерно-педагогический университет имени В.Г. Короленко»",
    "Сарапульский техникум машиностроения и информационных технологий"
]
ITEMS_PER_PAGE = 10

# ----------------- РАЙОНЫ УДМУРТИИ -----------------
DISTRICTS_UNIVERSITIES = {
    "Ижевск": [
        "БПОУ УР «Ижевский торгово-экономический техникум»",
        "БПОУ УР «Ижевский монтажный техникум»",
        "БПОУ «Ижевский агростроительный техникум»",
        "ЧПОО «Нефтяной техникум»",
        "БПОУ УР «Ижевский политехнический колледж»",
        "БПОУ УР «Ижевский промышленно-экономический колледж»",
        "БПОУ УР «Ижевский машиностроительный техникум им. С.Н. Борина»",
        "БПОУ УР «Радиомеханический техникум имени В.А. Шутова»",
        "АПОУ УР «Экономико-технологический колледж»",
        "АПОУ УР «Топливно-энергетический колледж»",
        "БПОУ УР «Ижевский техникум индустрии питания»",
        "КПОУ УР «Удмуртский республиканский колледж культуры»",
        "АНПОО «Международный Восточно-Европейский колледж»",
        "АПОУ УР «Техникум радиоэлектроники и информационных технологий им. А.В. Воскресенского»",
        "АПОУ УР «Республиканский медицинский колледж имени Героя Советского Союза Ф.А. Пушиной Министерства здравоохранения Удмуртской Республики»",
        "ПОЧУ «Ижевский техникум экономики, управления и права Удмуртпотребсоюза»",
        "АНПОО СПО «Ижевский финансово-юридический колледж»",
        "БПОУ УР «Удмуртский республиканский социально-педагогический колледж»",
        "АПОУ УР «Строительный техникум»",
        "ФГБОУ ВО «Ижевская государственная медицинская академия»",
        "КПОУ УР «Республиканский музыкальный колледж»",
        "ФГБОУ ВО «Приволжский государственный университет путей сообщения»",
        "БПОУ УР «Ижевский индустриальный техникум имени Евгения Фёдоровича Драгунова»",
        "Ижевский институт (филиал) ВГУЮ (РПА Минюста России)",
        "ФГБОУ ВО «Удмуртский государственный университет»",
        "БПОУ УР «Ижевский автотранспортный техникум»",
        "ФГБОУ ВО «Удмуртский государственный аграрный университет»",
        "ФГБОУ ВО «Ижевский государственный технический университет имени М.Т. Калашникова»",
        "Министерство юстиции Российской Федерации"
        "БПОУ УР «Ижевский автотранспортный техникум»",
    ],
    "Воткинск": [
        "БПОУ УР «Воткинский промышленный техникум»",
        "БПОУ УР «Воткинский музыкально-педагогический колледж имени П.И. Чайковского»",
        "Воткинский филиал АПОУ УР «Республиканский медицинский колледж имени Героя Советского Союза Ф.А. Пушиной Министерства здравоохранения Удмуртской Республики»",
        "БПОУ УР «Воткинский машиностроительный техникум имени В.Г.Садовникова»",
    ],
    "Глазов": [
        "АПОУ УР «Глазовский аграрно-промышленный техникум»",
        "БПОУ УР «Глазовский технический колледж»",
        "БПОУ УР «Глазовский политехнический колледж»",
        "Республиканский медицинский колледж имени Героя Советского Союза Ф.А. Пушиной, Глазовский филиал",
        "ФГБОУ ВО «Глазовский государственный инженерно-педагогический университет имени В. Г. Короленко»",
        "ФГБОУ ВО «Глазовский государственный инженерно-педагогический университет имени В. Г. Корененко»",
        "Глазовский инженерно-экономический институт(филиал) ФГБОУ ВО «ИжГТУ имени М.Т. Калашникова»",
        "Глазовский инженерно-экономический институт(филиал) ФГБОУ ВО «ИжГТУ имени М.Т. Калашникова»",
        "Глазовский инженерно-экономический институт(филиал) ФГБОУ ВО «ИжГТУ имени М.Т. Калашникова»",
        "ФГБОУ ВО «Глазовский государственный инженерно-педагогический университет имени В.Г. Короленко»",
    ],
    "Можга": [
        "БПОУ УР «Ижевский промышленно-экономический колледж» в г. Можга",
        "БПОУ УР «Можгинский педагогический колледж имени Т.К. Борисова»",
        "БПОУ УР «Можгинский агропромышленный колледж»",
    ],
    "Сарапул": [
        "БПОУ УР «Сарапульский политехнический техникум»",
        "БПОУ УР «Сарапульский многопрофильный колледж»",
        "БПОУ УР «Сарапульский колледж социально-педагогических технологий и сервиса»",
        "Сарапульский техникум машиностроения и информационных технологий",
    ],
    "Алнаши": ["БПОУ УР «Асановский аграрно-технический техникум»"],
    "Дебесы": ["БПОУ «Дебесский политехникум»"],
    "Игра": ["БПОУР «Игринский политехнический техникум»"],
    "Сюмси": ["БПОУ УР «Сюмсинский техникум лесного и сельского хозяйства»"],
    "Ува": ["БПОУ УР «Увинский профессиональный колледж»"],
    "Яр": ["БПОУ УР «Ярский политехникум»"],
}

INSTITUTION_TO_DISTRICT = {}
for _district, _unis in DISTRICTS_UNIVERSITIES.items():
    for _uni in _unis:
        INSTITUTION_TO_DISTRICT[_uni] = _district

# ----------------- ШАГИ АНКЕТЫ -----------------
STEPS = [
    "fio", "consent", "institution", "specialty", "study_group",
    "course", "form_of_study", "contacts",
    "employment_status", "target_contract", "experience",
    "practice_eval", "events", "resume_status",
    "interview_training", "special_status", "military",
    "maternity", "graduate", "post_plans", "help_needed"
]

STEP_TO_DB = {s: s for s in STEPS}

QUESTIONS = {
    "fio": "Пожалуйста, укажите ваши фамилию, имя и отчество полностью:",
    "consent": (
        "Я, {fio}, на основании статей 9, 11 Федерального закона от 27 июля 2006 г. N 152-ФЗ "
        "\"О персональных данных\" в целях моей профессиональной ориентации даю свое согласие "
        "казенному учреждению Удмуртской Республики «Республиканский центр занятости населения» "
        "на автоматизированную, а также без использования средств автоматизации обработку своих "
        "персональных данных, включая сбор, систематизацию, накопление, хранение, уточнение "
        "(обновление, изменение), использование, обезличивание, блокирование, уничтожение "
        "персональных данных о моих фамилии, имени, отчестве, номере телефона, адресе электронной почты.\n\n"
        "Настоящее согласие действует в течение 1 года с даты анкетирования.\n\n"
        "Пожалуйста, подтвердите согласие, нажав кнопку ниже:"
    ),
    "institution": "Выберите ваше учебное заведение из списка (используйте «далее» / «назад» для пролистывания):",
    "specialty": "Укажите вашу специальность обучения:",
    "study_group": "Укажите номер вашей учебной группы:",
    "course": "Выберите ваш курс обучения:",
    "form_of_study": "Выберите форму обучения:",
    "contacts": "Укажите контактные данные — телефон и/или e-mail (например: +79991234567 student@mail.ru). Можно указать оба контакта через пробел или запятую:",
    "employment_status": "Ваш статус занятости прямо сейчас. Выберите один вариант, указав его номер:",
    "target_contract": "Есть ли у вас заключённый договор о целевом обучении с работодателем?",
    "experience": "Есть ли у вас опыт работы или оплачиваемой стажировки по основной или близкой к ней специальности?",
    "practice_eval": "Как вы в целом оцениваете результаты своих производственных практик у работодателей по специальности обучения?",
    "events": "Участвовали ли вы в течение обучения в мероприятиях, которые помогают познакомиться с работодателями (ярмарки вакансий, дни карьеры, профтуры на предприятия, встречи с работодателями и т.д.)?",
    "resume_status": "Наличие резюме для поиска работы:",
    "interview_training": "Проходили ли вы занятия или тренинги по навыкам прохождения собеседования?",
    "special_status": "Есть ли у вас особый статус или жизненные обстоятельства?",
    "military": "Планируется ли в отношении вас призыв на военную службу в ближайшее время (после окончания текущего года обучения)?",
    "maternity": "Планируете ли вы уходить в отпуск по уходу за ребёнком в период обучения или сразу после окончания обучения (или продолжать уже начатый отпуск)?",
    "graduate": "Являетесь ли вы студентом выпускного курса (оканчиваете программу в текущем учебном году)?",
    "post_plans": "Ваши планы после выпуска. Выберите один или несколько вариантов, указав их номера через запятую (например: 1, 3):",
    "help_needed": "Какую помощь от Кадрового центра «Работа России» вы бы считали наиболее полезной? Выберите все подходящие варианты, указав их номера через запятую (например: 1, 2, 4):"
}

OPTIONS = {
    "consent": ["Да, я согласен(на)", "Нет, я не согласен(на)"],
    "course": ["1 курс", "2 курс", "3 курс", "4 курс", "5 курс"],
    "form_of_study": ["очная", "очно-заочная", "заочная"],
    "employment_status": [
        "Работаю по трудовому договору (в том числе по совместительству)",
        "Работаю по гражданско-правовому договору (договор подряда, услуг и т.п.)",
        "Являюсь самозанятым / ИП / учредителем юрлица",
        "Прохожу оплачиваемую стажировку / практику у работодателя",
        "Работаю временно (разовые подработки), не по специальности обучения",
        "Ничего из вышеперечисленного"
    ],
    "target_contract": ["да, договор о целевом обучении заключён", "нет, договора о целевом обучении нет"],
    "experience": [
        "да, есть опыт работы / оплачиваемой стажировки по основной или близкой специальности",
        "есть опыт работы только вне специальности обучения",
        "нет, опыта работы и оплачиваемых стажировок пока не было"
    ],
    "practice_eval": [
        "скорее доволен(льна) или полностью доволен(льна)",
        "скорее не доволен(на) / совсем не доволен(льна) результатами практик",
        "не проходил(а) производственную практику"
    ],
    "events": [
        "да, за последний год участвовал(а) хотя бы в одном таком мероприятии",
        "участвовал(а), но более года назад",
        "нет, ещё ни разу не участвовал(а)"
    ],
    "resume_status": [
        "есть актуальное резюме, которым я пользуюсь или готов(а) пользоваться",
        "резюме есть, но оно устарело / резюме нет, я его не составлял(а)"
    ],
    "interview_training": ["да, проходил(а) одно или несколько таких мероприятий", "пока не проходил(а)"],
    "special_status": [
        "да, имею группу инвалидности",
        "отношусь к категории детей-сирот и детей, оставшихся без попечения родителей",
        "планирую переезд в другой регион / страну после окончания обучения",
        "ничего из вышеперечисленного"
    ],
    "military": ["да, планируется призыв", "нет / не подлежу призыву / вопрос уже решён (служба пройдена и др.)"],
    "maternity": ["да, планирую", "пока не планирую"],
    "graduate": ["да, я учусь на выпускном курсе", "нет, я не на выпускном курсе"],
    "post_plans": [
        "У меня есть подписанный трудовой договор (или договор на целевое обучение)",
        "Есть устная договорённость с работодателем, но без подписанных документов",
        "Прохожу стажировку",
        "Планирую организовать своё дело (самозанятость / ИП / учредитель юрлица)",
        "Планирую продолжить обучение (магистратура / аспирантура и пр.)",
        "Сейчас ищу работу",
        "Пока нет планов"
    ],
    "help_needed": [
        "Подбор актуальных вакансий с учётом специальности",
        "Тренинги по составлению резюме, подготовке к собеседованиям, сопроводительных писем",
        "Профтур-экскурсии на предприятия",
        "Подбор оплачиваемой стажировки",
        "Подбор работодателя для практики",
        "Заключение договора с работодателем на целевое обучение",
        "Помощь с ЕЦП «Работа России»",
        "Другое (укажите)"
    ]
}

MULTI_STEPS = ["post_plans", "help_needed"]

MESSAGES = {
    "welcome": (
        "👋 Здравствуйте!\n\n"
        "Я задам несколько коротких вопросов о вашем обучении и занятости. "
        "Это займёт примерно 5–7 минут. По итогу анкетирования кадровый центр "
        "«Работа России» поможет Вам в прохождении тестирования на определение "
        "склонностей к профессиям, в составлении грамотного резюме, "
        "а также в подборе подходящих вакансий.\n\n"
        "Нажмите кнопку «Начать анкету», чтобы приступить."
    ),
    "invalid_contact": (
        "Не удалось распознать контакты. Пожалуйста, введите:\n\n"
        "• Номер телефона в формате +79991234567 или 89991234567\n"
        "и/или\n"
        "• Адрес электронной почты в формате example@mail.ru\n\n"
        "Можно указать оба контакта через пробел или запятую."
    ),
    "invalid_number": "Пожалуйста, введите номер от 1 до {}.",
    "invalid_multi": "Пожалуйста, укажите номера вариантов через запятую (например: 1, 3, 5). Проверьте, что номера от 1 до {}.",
    "no_data": "Пока нет собранных анкет для выгрузки.",
    "no_data_today": "Сегодня пока нет новых анкет для выгрузки.",
    "admin_only": "Эта команда доступна только администраторам.",
    "finished": "✅ Спасибо! Анкета заполнена.\n\nПредоставленная вами информация позволит нам детально проанализировать ситуацию и предложить оптимальное решение.\n\nЕсли хотите пройти анкету заново — нажмите кнопку ниже.",
    "already_finished": "✅ Вы уже заполнили анкету ранее. Если хотите пройти заново — нажмите кнопку ниже.",
}

# ----------------- КЛАВИАТУРЫ -----------------

def kb_start():
    return json.dumps({"one_time": False, "buttons": [[{"action": {"type": "text", "label": "Начать анкету"}, "color": "positive"}]]})

def kb_restart():
    return json.dumps({"one_time": False, "buttons": [[{"action": {"type": "text", "label": "🔄 Пройти заново"}, "color": "negative"}]]})

# ----------------- ОТПРАВКА И ВОПРОСЫ -----------------

def send_message(user_id, message, keyboard=None, attachment=None):
    try:
        params = {"peer_id": user_id, "message": message, "random_id": get_random_id()}
        if keyboard:
            params["keyboard"] = keyboard
        if attachment:
            params["attachment"] = attachment
        vk.messages.send(**params)
    except Exception as e:
        log.error(f"Ошибка отправки: {e}")

def ask_university_page(user_id, page):
    page = page if isinstance(page, int) else 0
    start = page * ITEMS_PER_PAGE
    end = min(start + ITEMS_PER_PAGE, len(UNIVERSITIES))
    items = UNIVERSITIES[start:end]
    list_text = format_numbered_list(items, start_from=start + 1, truncate=True)
    nav = []
    if page > 0:
        nav.append("«назад» — предыдущая страница")
    if end < len(UNIVERSITIES):
        nav.append("«далее» — следующая страница")
    nav_text = "\n".join(nav) if nav else ""
    message = f"{QUESTIONS['institution']}\n\n{list_text}"
    if nav_text:
        message += f"\n\n{nav_text}"
    message += "\n\nВведите номер вашего учебного заведения."
    send_message(user_id, message)

def ask_step(user_id, step_key, uni_page=0):
    uni_page = uni_page if isinstance(uni_page, int) else 0
    if step_key == "institution":
        ask_university_page(user_id, uni_page)
    elif step_key == "consent":
        conn = get_db()
        try:
            c = conn.cursor()
            c.execute("SELECT fio FROM answers WHERE user_id=%s", (user_id,))
            row = c.fetchone()
            fio_text = row[0] if row and row[0] else "[ФИО не указано]"
        finally:
            release_db(conn)
        message = QUESTIONS["consent"].format(fio=fio_text)
        opts = OPTIONS["consent"]
        list_text = format_numbered_list(opts, truncate=False)
        hint = "Напишите номер выбранного варианта (1 или 2)."
        send_message(user_id, f"{message}\n\n{list_text}\n\n{hint}")
    elif step_key in OPTIONS:
        opts = OPTIONS[step_key]
        list_text = format_numbered_list(opts, truncate=False)
        if step_key in MULTI_STEPS:
            hint = "Напишите номера выбранных вариантов через запятую (например: 1, 3)."
        else:
            hint = "Напишите номер выбранного варианта (например: 1)."
        message = f"{QUESTIONS[step_key]}\n\n{list_text}\n\n{hint}"
        send_message(user_id, message)
    else:
        send_message(user_id, QUESTIONS[step_key])

def advance_step(user_id, step_index):
    next_idx = step_index + 1
    if next_idx >= len(STEPS):
        conn = get_db()
        try:
            c = conn.cursor()
            c.execute("UPDATE answers SET created_at=%s WHERE user_id=%s", (datetime.now(), user_id))
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            release_db(conn)
        set_progress(user_id, next_idx, 0, 2)
        send_message(user_id, MESSAGES["finished"], kb_restart())
    else:
        set_progress(user_id, next_idx, 0, 1)
        ask_step(user_id, STEPS[next_idx])

# ----------------- ПАРСИНГ -----------------

def parse_single_number(text, max_val):
    text = text.strip()
    if text.isdigit():
        n = int(text)
        if 1 <= n <= max_val:
            return n
    return None

def parse_multi_numbers(text, max_val):
    try:
        parts = [p.strip() for p in text.split(",") if p.strip()]
        nums = [int(p) for p in parts]
        if any(n < 1 or n > max_val for n in nums):
            return None
        return nums
    except ValueError:
        return None

# ----------------- ПОДСЧЁТ БАЛЛОВ -----------------

def calculate_scores(row):
    scores = {}
    emp = (row.get("employment_status") or "").lower()
    if "трудовому договору" in emp:
        scores["employment"] = 0
    elif any(k in emp for k in ["гражданско-правовому", "самозанят", "стажировк", "временно"]):
        scores["employment"] = 0
    elif "ничего из вышеперечисленного" in emp:
        scores["employment"] = 1
    tc = (row.get("target_contract") or "").lower()
    if "да" in tc and "нет" not in tc:
        scores["target_contract"] = 0
    elif "нет" in tc:
        scores["target_contract"] = 2
    exp = (row.get("experience") or "").lower()
    if "да, есть опыт" in exp:
        scores["experience"] = 0
    elif "вне специальности" in exp:
        scores["experience"] = 3
    elif "нет, опыта" in exp:
        scores["experience"] = 3
    pe = (row.get("practice_eval") or "").lower()
    if "не доволен" in pe or "недоволен" in pe:
        scores["practice_eval"] = 2
    elif "доволен" in pe:
        scores["practice_eval"] = 0
    elif "не проходил" in pe:
        scores["practice_eval"] = 2
    ev = (row.get("events") or "").lower()
    if "за последний год" in ev:
        scores["events"] = 0
    elif "более года назад" in ev or "ни разу" in ev:
        scores["events"] = 2
    rs = (row.get("resume_status") or "").lower()
    if "актуальное" in rs:
        scores["resume"] = 0
    elif "устарело" in rs or "не составлял" in rs:
        scores["resume"] = 1
    it = (row.get("interview_training") or "").lower()
    if "да" in it and "не проходил" not in it:
        scores["interview"] = 1
    elif "не проходил" in it:
        scores["interview"] = 0
    ss = (row.get("special_status") or "").lower()
    if "ничего из вышеперечисленного" in ss:
        scores["special_status"] = 0
    elif ss:
        scores["special_status"] = 1
    mil = (row.get("military") or "").lower()
    if "да, планируется" in mil:
        scores["military"] = 1
    elif "нет" in mil or "не подлежу" in mil:
        scores["military"] = 0
    total = sum(v for v in scores.values())
    return scores, total

# ----------------- ВЫГРУЗКА -----------------

EXPORT_HEADERS = {
    "user_id": "ID пользователя", "fio": "ФИО", "institution": "Учебное заведение",
    "specialty": "Специальность", "study_group": "Учебная группа", "course": "Курс",
    "form_of_study": "Форма обучения", "contacts": "Контакты",
    "employment_status": "Статус занятости", "target_contract": "Целевой договор",
    "experience": "Опыт работы", "practice_eval": "Оценка практик",
    "events": "Участие в мероприятиях", "resume_status": "Наличие резюме",
    "interview_training": "Тренинги по собеседованию", "special_status": "Особый статус",
    "military": "Призыв на военную службу", "maternity": "Отпуск по уходу за ребёнком",
    "graduate": "Выпускной курс", "post_plans": "Планы после выпуска",
    "help_needed": "Нужная помощь", "created_at": "Дата заполнения",
}

def export_to_table(admin_id, today_only=False):
    conn = get_db()
    try:
        c = conn.cursor(cursor_factory=RealDictCursor)
        if today_only:
            today = date.today()
            c.execute("SELECT a.* FROM answers a LEFT JOIN progress p ON a.user_id = p.user_id WHERE a.created_at::date = %s ORDER BY p.started_at ASC NULLS FIRST", (today,))
        else:
            c.execute("SELECT a.* FROM answers a LEFT JOIN progress p ON a.user_id = p.user_id ORDER BY p.started_at ASC NULLS FIRST")
        rows = c.fetchall()
    finally:
        release_db(conn)

    if not rows:
        send_message(admin_id, MESSAGES["no_data_today"] if today_only else MESSAGES["no_data"])
        return

    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Анкеты"
    cols = list(rows[0].keys())
    header_font = Font(bold=True)
    for col_idx, col_name in enumerate(cols, start=1):
        cell = ws1.cell(row=1, column=col_idx, value=EXPORT_HEADERS.get(col_name, col_name))
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for row_idx, r in enumerate(rows, start=2):
        for col_idx, col_name in enumerate(cols, start=1):
            val = r[col_name]
            if col_name == "created_at" and val is not None:
                val = val.strftime("%Y-%m-%d %H:%M")
            ws1.cell(row=row_idx, column=col_idx, value=val if val is not None else "")
    for col in ws1.columns:
        ws1.column_dimensions[col[0].column_letter].width = 25

    bold_font = Font(bold=True)
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    district_headers = ["ФИО","Учебное заведение","Контакты","Статус занятости","Целевой договор","Опыт работы","Оценка практик","Мероприятия","Резюме","Собеседование","Особый статус","Военный призыв","Сумма баллов"]
    score_keys = ["employment","target_contract","experience","practice_eval","events","resume","interview","special_status","military"]

    def write_score_sheet(workbook, sheet_name, sheet_rows):
        ws = workbook.create_sheet(title=sheet_name)
        for col_idx, h in enumerate(district_headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        sorted_rows = sorted(sheet_rows, key=lambda r: r.get("created_at") or datetime.min)
        row_idx = 2
        for r in sorted_rows:
            scores, total = calculate_scores(dict(r))
            ws.cell(row=row_idx, column=1, value=r.get("fio") or "")
            ws.cell(row=row_idx, column=2, value=r.get("institution") or "")
            ws.cell(row=row_idx, column=3, value=r.get("contacts") or "")
            for i, key in enumerate(score_keys, start=4):
                val = scores.get(key)
                ws.cell(row=row_idx, column=i, value=val if val is not None else "")
            total_cell = ws.cell(row=row_idx, column=13, value=total)
            total_cell.font = bold_font
            total_cell.fill = total_fill
            row_idx += 1
        for col_idx in range(1, len(district_headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = 25

    for district_name in DISTRICTS_UNIVERSITIES:
        district_rows = [r for r in rows if INSTITUTION_TO_DISTRICT.get(r.get("institution")) == district_name]
        if district_rows:
            write_score_sheet(wb, district_name, district_rows)
    other_rows = [r for r in rows if r.get("institution") and INSTITUTION_TO_DISTRICT.get(r.get("institution")) is None]
    if other_rows:
        write_score_sheet(wb, "Прочие", other_rows)

    fname = tempfile.mktemp(suffix=".xlsx")
    wb.save(fname)
    if not os.path.exists(fname) or os.path.getsize(fname) == 0:
        send_message(admin_id, "❌ Не удалось создать файл выгрузки.")
        try: os.remove(fname)
        except: pass
        return

    try:
        upload_server = vk.docs.getMessagesUploadServer(type='doc', peer_id=admin_id)
        upload_url = upload_server['upload_url']
        with open(fname, "rb") as f:
            resp = requests.post(upload_url, files={"file": ("survey_export.xlsx", f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        if not resp.content:
            raise RuntimeError("VK вернул пустой ответ.")
        result = resp.json()
        if "file" not in result or not result["file"]:
            raise RuntimeError(f"VK не принял файл: {result}")
        file_data = result["file"]
        file_title = f"Выгрузка за {date.today().strftime('%d.%m.%Y')}" if today_only else "Выгрузка анкет"
        saved = vk.docs.save(file=file_data, title=file_title)
        if isinstance(saved, dict) and "doc" in saved:
            d = saved["doc"]
        elif isinstance(saved, dict) and "docs" in saved and len(saved["docs"]) > 0:
            d = saved["docs"][0]
        else:
            raise RuntimeError(f"Неожиданный ответ docs.save: {saved}")
        attachment = f"doc{d['owner_id']}_{d['id']}"
        sheet_list = []
        for district_name in DISTRICTS_UNIVERSITIES:
            count = sum(1 for r in rows if INSTITUTION_TO_DISTRICT.get(r.get("institution")) == district_name)
            if count:
                sheet_list.append(f"  • «{district_name}» — {count} чел.")
        if other_rows:
            sheet_list.append(f"  • «Прочие» — {len(other_rows)} чел.")
        sheets_text = "\n".join(sheet_list) if sheet_list else ""
        header = f"📊 Выгрузка анкет за сегодня ({date.today().strftime('%d.%m.%Y')}):\n\n" if today_only else "📊 Вот полная выгрузка анкет:\n\n"
        send_message(admin_id, f"{header}• Лист «Анкеты» — полные ответы\n• Листы по районам — ФИО, баллы:\n\n{sheets_text}\n\nВсего анкет: {len(rows)}", attachment=attachment)
    except Exception as e:
        log.error(f"Ошибка загрузки .xlsx: {e}")
        send_message(admin_id, f"❌ Не удалось отправить Excel: {e}\n\nПроверьте права токена.")
    finally:
        try: os.remove(fname)
        except: pass

# ----------------- ОБРАБОТКА НЕПРОЧИТАННЫХ -----------------

def process_unread_messages():
    processed = 0
    skipped = 0
    try:
        result = vk.messages.getConversations(filter='unread', count=100, extended=0)
    except Exception as e:
        log.error(f"Ошибка getConversations: {e}")
        return
    items = result.get('items', [])
    if not items:
        log.info("Непрочитанных сообщений нет.")
        return
    for conv in items:
        conv_info = conv.get('conversation', {})
        peer_id = conv_info.get('peer', {}).get('id')
        unread_count = conv_info.get('unread_count', 0)
        if not peer_id or unread_count == 0:
            continue
        try:
            history = vk.messages.getHistory(peer_id=peer_id, count=min(unread_count + 5, 200), extended=0)
        except Exception as e:
            log.error(f"Ошибка getHistory для peer_id={peer_id}: {e}")
            continue
        messages = history.get('items', [])
        if not messages:
            continue
        last_msg = messages[0]
        if last_msg.get('from_id', 0) < 0:
            skipped += 1
            try: vk.messages.markAsRead(peer_id=peer_id)
            except: pass
            continue
        incoming = [m for m in messages if m.get('from_id', 0) > 0]
        incoming.reverse()
        if len(incoming) > unread_count:
            incoming = incoming[-unread_count:]
        for msg in incoming:
            text = msg.get('text', '').strip()
            if text:
                try:
                    handle_message(peer_id, text)
                    processed += 1
                except Exception as e:
                    log.error(f"Ошибка обработки от {peer_id}: {e}")
        try: vk.messages.markAsRead(peer_id=peer_id)
        except: pass
    log.info(f"Обработано непрочитанных: {processed}, пропущено: {skipped}")

# ----------------- ОСНОВНАЯ ЛОГИКА -----------------

def handle_message(user_id, text):
    if not text or not user_id:
        return
    text_lower = text.lower()

    if text_lower in ["/export", "/выгрузить"]:
        if user_id in ADMIN_IDS:
            export_to_table(user_id, today_only=False)
        else:
            send_message(user_id, MESSAGES["admin_only"])
        return

    if text_lower in ["/export today", "/выгрузить сегодня", "/выгрузить_сегодня"]:
        if user_id in ADMIN_IDS:
            export_to_table(user_id, today_only=True)
        else:
            send_message(user_id, MESSAGES["admin_only"])
        return

    if text_lower == "/restart":
        set_progress(user_id, 0, 0, 0)
        send_message(user_id, "Анкета сброшена. Нажмите «Начать анкету».", kb_start())
        return

    if text == "🔄 Пройти заново":
        set_progress(user_id, 0, 0, 0)
        send_message(user_id, MESSAGES["welcome"], kb_start())
        return

    step_index, uni_page, started = get_progress(user_id)
    step_index = step_index if isinstance(step_index, int) else 0
    uni_page = uni_page if isinstance(uni_page, int) else 0
    started = started if isinstance(started, int) else 0

    if started == 0:
        if text == "Начать анкету":
            set_progress(user_id, 0, 0, 1)
            ask_step(user_id, STEPS[0])
        else:
            send_message(user_id, MESSAGES["welcome"], kb_start())
        return

    if started == 2 or step_index >= len(STEPS):
        send_message(user_id, MESSAGES["already_finished"], kb_restart())
        return

    step_key = STEPS[step_index]

    if step_key == "fio":
        ok, value = validate_fio(text)
        if ok:
            save_answer(user_id, "fio", value)
            advance_step(user_id, step_index)
        else:
            send_message(user_id, value)
        return

    if step_key == "institution":
        if text.lower() in ["далее", ">", "следующий"]:
            max_page = (len(UNIVERSITIES) - 1) // ITEMS_PER_PAGE
            if uni_page < max_page:
                set_progress(user_id, step_index, uni_page + 1, 1)
                ask_university_page(user_id, uni_page + 1)
            else:
                send_message(user_id, "Это последняя страница.")
                ask_university_page(user_id, uni_page)
            return
        elif text.lower() in ["назад", "<", "←"]:
            if uni_page > 0:
                set_progress(user_id, step_index, uni_page - 1, 1)
                ask_university_page(user_id, uni_page - 1)
            else:
                send_message(user_id, "Это первая страница.")
                ask_university_page(user_id, uni_page)
            return
        if text.isdigit():
            idx = int(text) - 1
            if 0 <= idx < len(UNIVERSITIES):
                save_answer(user_id, "institution", UNIVERSITIES[idx])
                advance_step(user_id, step_index)
                return
        send_message(user_id, "Пожалуйста, введите номер учебного заведения из списка или используйте «далее» / «назад».")
        ask_university_page(user_id, uni_page)
        return

    if step_key == "contacts":
        ok, value = validate_contacts(text)
        if ok:
            save_answer(user_id, "contacts", value)
            advance_step(user_id, step_index)
        else:
            send_message(user_id, MESSAGES["invalid_contact"])
        return

    if step_key in OPTIONS:
        opts = OPTIONS[step_key]
        if step_key == "consent":
            n = parse_single_number(text, len(opts))
            if n is None:
                send_message(user_id, MESSAGES["invalid_number"].format(len(opts)))
                return
            is_consent = (n == 1)
            conn = get_db()
            try:
                c = conn.cursor()
                c.execute("INSERT INTO answers (user_id) VALUES (%s) ON CONFLICT (user_id) DO NOTHING", (user_id,))
                c.execute("UPDATE answers SET consent_status=%s WHERE user_id=%s", (is_consent, user_id))
                conn.commit()
            except Exception:
                conn.rollback()
                raise
            finally:
                release_db(conn)
            advance_step(user_id, step_index)
            return

        if step_key in MULTI_STEPS:
            nums = parse_multi_numbers(text, len(opts))
            if nums is None:
                send_message(user_id, MESSAGES["invalid_multi"].format(len(opts)))
                return
            label = "; ".join(opts[n - 1] for n in nums)
            save_answer(user_id, STEP_TO_DB[step_key], label)
        else:
            n = parse_single_number(text, len(opts))
            if n is None:
                send_message(user_id, MESSAGES["invalid_number"].format(len(opts)))
                return
            save_answer(user_id, STEP_TO_DB[step_key], opts[n - 1])
        advance_step(user_id, step_index)
        return

    if step_key not in ["post_plans", "help_needed"]:
        if len(text) < 2:
            send_message(user_id, "Пожалуйста, введите более развёрнутый ответ.")
            return

    save_answer(user_id, STEP_TO_DB[step_key], text)
    advance_step(user_id, step_index)

# ==================== FASTAPI ====================

app = FastAPI()
bot_start_time = time.time()
_executor = None

@app.on_event("startup")
async def startup_event():
    global _executor
    from concurrent.futures import ThreadPoolExecutor
    _executor = ThreadPoolExecutor(max_workers=20)
    init_db_pool()
    init_db()
    log.info(f"Бот запущен. Время старта: {time.strftime('%Y-%m-%d %H:%M:%S', time.localtime(bot_start_time))}")
    log.info("Проверяю непрочитанные сообщения...")
    threading.Thread(target=process_unread_messages, daemon=True).start()

@app.post("/")
async def vk_callback(request: Request):
    try:
        data = await request.json()
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid JSON")

    event_type = data.get("type")
    if not event_type:
        raise HTTPException(status_code=400, detail="No event type")

    if event_type == "confirmation":
        return PlainTextResponse(CALLBACK_CONFIRM_TOKEN)

    if event_type == "message_new":
        obj = data.get("object", {})
        msg = obj.get("message", {})
        user_id = msg.get("from_id") or msg.get("peer_id")
        text = (msg.get("text") or "").strip()
        msg_time = msg.get("date")

        if not user_id or not text:
            return PlainTextResponse("ok")

        if msg_time and msg_time < bot_start_time:
            try:
                vk.messages.markAsRead(peer_id=user_id)
            except Exception:
                pass
            return PlainTextResponse("ok")

        loop = asyncio.get_running_loop()  # ← было get_event_loop()
        loop.run_in_executor(_executor, handle_message, user_id, text)
        return PlainTextResponse("ok")

    return PlainTextResponse("ok")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=int(os.getenv("PORT", "8000")))
